"""
02_association_rules.py  -  Association rules between skills and job titles.

Each posting is one transaction; each skill is one item.
Rules are counted directly (no Apriori library):
  * skill            -> job title   (antecedent size 1)
  * skill + skill    -> job title   (antecedent size 2)
  * skill -> skill   inside one job title
Metrics:  Support    = P(antecedent and consequent)
          Confidence = P(consequent | antecedent)
          Lift       = Confidence / P(consequent)        (1 = unrelated)

Input : data/processed/jobs_clean.csv
Output: data/processed/skill_association_rules.xlsx
Run   : python 02_association_rules.py
"""
import itertools
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
PROC = ROOT / "data" / "processed"
MIN_SUPPORT = 0.001          # keep skill->job rules seen in >= 0.1% of all postings
WITHIN_MIN_SUPPORT = 0.10    # within a title: pair seen in >= 10% of that title's postings
WITHIN_MIN_CONF = 0.60


def skill_to_job_rules(df):
    n = len(df)
    job_count = df["Job_Title"].value_counts()
    c1, c2, cj = Counter(), Counter(), Counter()
    for job, skills in zip(df["Job_Title"], df["sk"]):
        s = sorted(skills)
        for x in s:
            c1[x] += 1
            cj[(x, job)] += 1
        for p in itertools.combinations(s, 2):
            c2[p] += 1
            cj[(p, job)] += 1
    rows = []
    for (a, job), v in cj.items():
        ant = (a,) if isinstance(a, str) else a
        ca = c1[a] if isinstance(a, str) else c2[a]
        conf = v / ca
        rows.append((" + ".join(ant), len(ant), job, v / n, conf, conf / (job_count[job] / n), v, ca))
    r = pd.DataFrame(rows, columns=["Antecedent", "Size", "Consequent", "Support", "Confidence", "Lift",
                                    "Count", "AntCount"])
    return r[r["Support"] >= MIN_SUPPORT]


def within_job_rules(df):
    out = []
    for job, g in df.groupby("Job_Title"):
        n = len(g)
        s1, s2 = Counter(), Counter()
        for skills in g["sk"]:
            for x in skills:
                s1[x] += 1
            for p in itertools.combinations(sorted(skills), 2):
                s2[p] += 1
        for (a, b), v in s2.items():
            for x, y in ((a, b), (b, a)):
                conf = v / s1[x]
                out.append((job, x, y, v / n, conf, conf / (s1[y] / n), v))
    r = pd.DataFrame(out, columns=["Job", "Antecedent", "Consequent", "Support", "Confidence", "Lift", "Count"])
    return r[(r["Support"] >= WITHIN_MIN_SUPPORT) & (r["Confidence"] >= WITHIN_MIN_CONF)]


def job_similarity(df):
    sets = {j: set(s for l in g["sk"] for s in l) for j, g in df.groupby("Job_Title")}
    rows = []
    for a, b in itertools.combinations(sorted(sets), 2):
        i, u = sets[a] & sets[b], sets[a] | sets[b]
        rows.append((a, b, len(i) / len(u), len(i), ", ".join(sorted(i))))
    return pd.DataFrame(rows, columns=["Job_A", "Job_B", "Jaccard", "N_Shared", "Shared_Skills"]).sort_values(
        "Jaccard", ascending=False)


def write_sheet(wb, name, d):
    ws = wb.create_sheet(name)
    for j, c in enumerate(d.columns, 1):
        x = ws.cell(1, j, c)
        x.font = Font(name="Arial", bold=True, color="FFFFFF")
        x.fill = PatternFill("solid", fgColor="305496")
    for i, row in enumerate(d.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if hasattr(v, "item"):
                v = v.item()
            c = ws.cell(i, j, v)
            c.font = Font(name="Arial", size=10)
            col = d.columns[j - 1]
            if col in ("Support", "Confidence"):
                c.number_format = "0.0%"
            elif col in ("Lift", "Jaccard"):
                c.number_format = "0.00"
    for j, c in enumerate(d.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(len(str(c)), *[len(str(v)) for v in d[c].head(100)]) + 2, 70)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    df = pd.read_csv(PROC / "jobs_clean.csv", usecols=["Job_ID", "Job_Title", "Skills"])
    df["sk"] = df["Skills"].str.split(", ")

    r1 = skill_to_job_rules(df)
    r2 = within_job_rules(df)
    sim = job_similarity(df)

    wb = Workbook()
    wb.remove(wb.active)
    s1 = r1[r1["Size"] == 1].sort_values(["Confidence", "Lift", "Support"], ascending=False)
    pairs = r1[(r1["Size"] == 2) & (r1["Confidence"] >= 0.9)].sort_values(["Support", "Lift"], ascending=False).head(300)
    write_sheet(wb, "Skill_to_Job", s1[["Antecedent", "Consequent", "Support", "Confidence", "Lift", "Count", "AntCount"]])
    write_sheet(wb, "Pair_to_Job", pairs[["Antecedent", "Consequent", "Support", "Confidence", "Lift", "Count"]])
    write_sheet(wb, "Within_Job_Rules", r2.sort_values(["Job", "Lift"], ascending=[True, False]))
    write_sheet(wb, "Job_Similarity", sim)
    out = PROC / "skill_association_rules.xlsx"
    wb.save(out)

    ex = r1[(r1["Antecedent"].isin(["Excel + Python", "Python + Excel"])) & (r1["Consequent"] == "Data Analyst")]
    print("example {Excel, Python} -> Data Analyst:")
    print(ex[["Support", "Confidence", "Lift", "Count"]].round(4).to_string(index=False))
    print("written to", out)


if __name__ == "__main__":
    main()
